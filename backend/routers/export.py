from __future__ import annotations
import io

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from models import ExportRequest

router = APIRouter(tags=["export"])

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

COLUMNS = [
    ("Equipment Type", "equipment_type"),
    ("Brand", "brand"),
    ("Model", "model"),
    ("Installation Date", "installation_date"),
    ("Last Maintenance", "last_maintenance_date"),
    ("Next Maintenance", "next_maintenance_date"),
    ("Interval (months)", "maintenance_interval_months"),
    ("Manual Source", "manual_source"),
    ("Notes", "gemini_reasoning"),
]


@router.post("/export-xlsx")
async def export_xlsx(body: ExportRequest):
    wb = Workbook()
    ws = wb.active
    ws.title = "Maintenance Schedule"

    # Write headers
    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    # Write data rows
    for row_idx, item in enumerate(body.equipment, start=2):
        data = item.model_dump(mode="json")
        fill = ALT_FILL if row_idx % 2 == 0 else None
        for col_idx, (_, field) in enumerate(COLUMNS, start=1):
            value = data.get(field)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill:
                cell.fill = fill

    # Auto-size columns
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=maintenance_schedule.xlsx"},
    )
